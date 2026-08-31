"""
Reading the plant's own Excel register back into the system.

WHY THIS EXISTS AT ALL

Greenlam already keeps a breakdown register in Excel, and will keep keeping it
during the pilot — that is not a habit to fight, it is the source of truth the
plant currently trusts. Two consequences follow:

  1. History. The system launches knowing nothing. The register holds years.
     Importing it is the difference between a Pareto built on three weeks and a
     Pareto built on three years, and only one of those changes a decision.

  2. The parallel period. During a pilot the register keeps being filled in by
     people who are not yet using the app. Without an import, the two diverge
     immediately and the app looks wrong — which is how pilots die.

THE SHAPE IS THE EXPORT'S SHAPE, DELIBERATELY

`export/build_workbook.py` writes a BD Tracker sheet with a fixed set of
headers, quirks included — the double space in "Problem observed after how many
sheet  sanding" is load-bearing, because leadership's pivot tables reference it.
This importer reads exactly that shape. So the loop closes: export, edit in
Excel where people are fastest, upload again. A plant will do that whether or
not it is supported; supporting it is what stops it happening by email.

FOUR RULES THIS FILE ENFORCES

  1. DRY RUN FIRST, ALWAYS. `plan()` never writes. The caller decides whether to
     apply what it describes. Silently ingesting the wrong file is the most
     damaging thing this feature could do, and "are you sure" is not a defence —
     showing the reader exactly which 12 rows would change is.

  2. ROW NUMBERS ON EVERY PROBLEM. "Invalid file" is useless against a
     4,000-row register. Every issue names the sheet row as Excel numbers it, so
     the person can go and look at it.

  3. NEVER INVENT. A row with an unknown machine code is REPORTED, not guessed
     into the nearest match. A typo silently mapped to the wrong press corrupts
     exactly the analysis this system exists to produce.

  4. IDEMPOTENT. Each row gets a deterministic key from its own content. Import
     the same file twice and the second pass creates nothing — which matters
     because "did that upload work?" is answered by uploading again.
"""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from typing import Any, Literal

from openpyxl import load_workbook

# The BD Tracker headers, verbatim, including the double space. Kept in sync
# with export/build_workbook.py — see the test that asserts they match.
BD_HEADERS = [
    "Sr No",
    "Date",
    "Month",
    "Section",
    "Machine / Area",
    "Issue Category",
    "BD (Min)",
    "BD (Hrs)",
    "BD Details",
    "Immediate Correction(Filled by maint.)",
    "Root Cause (Why-Why) (Filled by maint.)",
    "Preventive Action(Filled by Maint)",
    "Problem observed after how many sheet  sanding",
]

# Columns without which a row cannot become a ticket at all.
REQUIRED = ("Date", "Machine / Area", "BD (Min)")

SHEET_CANDIDATES = ("BD Tracker", "Breakdown", "Sheet1")

Severity = Literal["error", "warning"]
Action = Literal["create", "update", "skip"]


@dataclass(frozen=True)
class RowIssue:
    row: int
    column: str
    message: str
    severity: Severity


@dataclass
class PlannedRow:
    row: int
    action: Action
    machine_code: str
    raised_at: datetime
    minutes: int
    description: str
    category: str | None
    immediate_correction: str | None
    root_cause: str | None
    preventive_action: str | None
    sheets_after_sanding: str | None
    import_key: str
    reason: str | None = None


@dataclass
class ImportPlan:
    file_hash: str
    sheet_name: str
    rows_read: int = 0
    planned: list[PlannedRow] = field(default_factory=list)
    issues: list[RowIssue] = field(default_factory=list)
    missing_columns: list[str] = field(default_factory=list)
    unknown_machines: list[str] = field(default_factory=list)

    @property
    def creates(self) -> int:
        return sum(1 for p in self.planned if p.action == "create")

    @property
    def updates(self) -> int:
        return sum(1 for p in self.planned if p.action == "update")

    @property
    def skips(self) -> int:
        return sum(1 for p in self.planned if p.action == "skip")

    @property
    def errors(self) -> int:
        return sum(1 for i in self.issues if i.severity == "error")

    @property
    def can_commit(self) -> bool:
        """A plan with no usable rows must not be committable.

        Not the same test as "no errors" — a 4,000-row register with 12 bad
        rows is perfectly importable, and refusing the whole file over 12 rows
        would send the user back to Excel to fix data the system could simply
        report. What is never acceptable is committing a file whose headers did
        not match, because that means every row was parsed against the wrong
        columns and the "successes" are nonsense.
        """
        return not self.missing_columns and (self.creates + self.updates) > 0


def file_digest(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _row_key(machine_code: str, when: datetime, minutes: int, description: str) -> str:
    """A row's identity, derived from its own content.

    The register has a "Sr No" column, but it is per-file and gets renumbered
    every time someone sorts the sheet — using it as an identity would make a
    re-sorted file look like an entirely new set of breakdowns. Machine, day,
    duration and the first words of the description are what actually identify
    an event to a human reading the register, so they are what identifies it
    here.
    """
    basis = "|".join(
        [
            machine_code.strip().upper(),
            when.date().isoformat(),
            str(minutes),
            re.sub(r"\s+", " ", description.strip().lower())[:80],
        ]
    )
    return hashlib.sha256(basis.encode()).hexdigest()[:32]


def _clean(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _coerce_date(value: Any) -> datetime | None:
    """Excel dates arrive as datetimes, dates, serial numbers or free text.

    openpyxl converts most real date cells for us. The two that reach here raw
    are a cell formatted as General (a float serial) and a cell someone typed
    as text. Both are common in a register maintained by many hands over years.
    """
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    if isinstance(value, (int, float)) and value > 0:
        # Excel's 1900 epoch, including its deliberate leap-year bug: serial 60
        # is the non-existent 29 Feb 1900, so anything after it is off by one.
        base = datetime(1899, 12, 30, tzinfo=UTC)
        return base + timedelta(days=float(value))
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=UTC)
        except ValueError:
            continue
    return None


def _coerce_minutes(minutes_cell: Any, hours_cell: Any) -> tuple[int | None, str | None]:
    """Downtime, in minutes, and a warning if the two columns disagree.

    The register carries both BD (Min) and BD (Hrs) and they are maintained by
    hand, so they drift. Minutes wins — it is the finer unit and the one people
    actually type — but a disagreement is worth surfacing, because a row where
    hours says 8 and minutes says 30 is a data-entry error somebody should look
    at rather than something to silently resolve.
    """
    minutes: int | None = None
    if isinstance(minutes_cell, (int, float)):
        minutes = int(round(float(minutes_cell)))
    else:
        text = _clean(minutes_cell)
        if text:
            try:
                minutes = int(round(float(text.replace(",", ""))))
            except ValueError:
                minutes = None

    hours: float | None = None
    if isinstance(hours_cell, (int, float)):
        hours = float(hours_cell)
    else:
        text = _clean(hours_cell)
        if text:
            try:
                hours = float(text)
            except ValueError:
                hours = None

    if minutes is None and hours is not None:
        return int(round(hours * 60)), None
    if minutes is not None and hours is not None and hours > 0:
        implied = hours * 60
        if implied > 0 and abs(implied - minutes) > max(5.0, 0.1 * implied):
            return minutes, (
                f"BD (Min) says {minutes} but BD (Hrs) says {hours:g} "
                f"({implied:.0f} min). Using minutes."
            )
    return minutes, None


def _pick_sheet(workbook: Any) -> Any:
    for name in SHEET_CANDIDATES:
        if name in workbook.sheetnames:
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def _header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
    """Map header text to column index.

    By NAME, not by position. Someone inserting a column in the middle of the
    register is the single most likely edit between exports, and a positional
    reader would silently shift every field one to the right — writing machine
    codes into the category column and never raising an error.

    Whitespace is normalised for matching only. The canonical names keep their
    quirks, including that double space, because those are what gets written
    back out.
    """
    normalise = lambda s: re.sub(r"\s+", " ", str(s).strip().lower())  # noqa: E731
    lookup = {normalise(h): i for i, h in enumerate(header_row) if h is not None}
    return {
        canonical: lookup[normalise(canonical)]
        for canonical in BD_HEADERS
        if normalise(canonical) in lookup
    }


def plan(
    payload: bytes,
    *,
    known_machines: set[str],
    existing_keys: set[str],
    known_categories: set[str] | None = None,
    max_rows: int = 20_000,
) -> ImportPlan:
    """Read a workbook and describe what importing it WOULD do. Writes nothing.

    `known_machines` and `existing_keys` are passed in rather than queried here
    so this function stays pure: it can be unit-tested against a fixture
    workbook with no database, which is what makes the parsing rules — Excel's
    leap-year bug, the minutes/hours disagreement, header reordering — actually
    covered by tests instead of merely believed.
    """
    # read_only keeps a 4,000-row register off the heap as a cell-object graph.
    # The buffer is closed explicitly alongside the workbook: openpyxl's
    # read-only mode holds the underlying zip open, and letting it be collected
    # raises "I/O operation on closed file" out of a __del__ where nothing can
    # handle it.
    buffer = io.BytesIO(payload)
    workbook = load_workbook(buffer, read_only=True, data_only=True)
    try:
        sheet = _pick_sheet(workbook)
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration:
            return ImportPlan(
                file_digest(payload), sheet.title, missing_columns=list(BD_HEADERS)
            )

        columns = _header_map(header_row)
        result = ImportPlan(file_digest(payload), sheet.title)
        result.missing_columns = [c for c in REQUIRED if c not in columns]
        if result.missing_columns:
            # Parsing further would produce confident nonsense.
            return result

        def cell(row: tuple[Any, ...], name: str) -> Any:
            index = columns.get(name)
            if index is None or index >= len(row):
                return None
            return row[index]

        seen_in_file: set[str] = set()
        unknown: set[str] = set()

        for offset, row in enumerate(rows):
            excel_row = offset + 2  # header is row 1, and Excel counts from 1
            if excel_row - 1 > max_rows:
                result.issues.append(
                    RowIssue(
                        excel_row,
                        "",
                        f"Stopped after {max_rows} rows. Split the file and upload again.",
                        "error",
                    )
                )
                break
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            result.rows_read += 1

            machine_code = _clean(cell(row, "Machine / Area"))
            when = _coerce_date(cell(row, "Date"))
            minutes, minute_warning = _coerce_minutes(
                cell(row, "BD (Min)"), cell(row, "BD (Hrs)")
            )
            description = _clean(cell(row, "BD Details")) or ""

            if not machine_code:
                result.issues.append(
                    RowIssue(excel_row, "Machine / Area", "No machine code.", "error")
                )
                continue
            if when is None:
                result.issues.append(
                    RowIssue(
                        excel_row,
                        "Date",
                        f"Could not read {cell(row, 'Date')!r} as a date.",
                        "error",
                    )
                )
                continue
            if minutes is None:
                result.issues.append(
                    RowIssue(excel_row, "BD (Min)", "No downtime recorded.", "error")
                )
                continue
            if minutes < 0:
                result.issues.append(
                    RowIssue(
                        excel_row, "BD (Min)", "Downtime cannot be negative.", "error"
                    )
                )
                continue
            if when > datetime.now(UTC) + timedelta(days=1):
                result.issues.append(
                    RowIssue(
                        excel_row,
                        "Date",
                        f"{when.date().isoformat()} is in the future.",
                        "error",
                    )
                )
                continue

            normalised = machine_code.strip().upper()
            if normalised not in known_machines:
                # Rule 3: report, never guess. A fuzzy match here would quietly
                # attribute a breakdown to the wrong press.
                unknown.add(machine_code)
                result.issues.append(
                    RowIssue(
                        excel_row,
                        "Machine / Area",
                        f"{machine_code!r} is not in the machine master. "
                        "Add it under masters, or correct the spelling.",
                        "error",
                    )
                )
                continue

            if minute_warning:
                result.issues.append(
                    RowIssue(excel_row, "BD (Min)", minute_warning, "warning")
                )

            category = _clean(cell(row, "Issue Category"))
            if category and known_categories and category not in known_categories:
                result.issues.append(
                    RowIssue(
                        excel_row,
                        "Issue Category",
                        f"{category!r} is not a known category — importing without one.",
                        "warning",
                    )
                )
                category = None

            key = _row_key(normalised, when, minutes, description)

            if key in seen_in_file:
                result.planned.append(
                    PlannedRow(
                        excel_row,
                        "skip",
                        machine_code,
                        when,
                        minutes,
                        description,
                        category,
                        None,
                        None,
                        None,
                        None,
                        key,
                        reason="Duplicate of an earlier row in this file.",
                    )
                )
                continue
            seen_in_file.add(key)

            action: Action = "update" if key in existing_keys else "create"
            result.planned.append(
                PlannedRow(
                    row=excel_row,
                    action=action,
                    machine_code=machine_code,
                    raised_at=when,
                    minutes=minutes,
                    description=description or f"Imported from register, row {excel_row}",
                    category=category,
                    immediate_correction=_clean(
                        cell(row, "Immediate Correction(Filled by maint.)")
                    ),
                    root_cause=_clean(
                        cell(row, "Root Cause (Why-Why) (Filled by maint.)")
                    ),
                    preventive_action=_clean(
                        cell(row, "Preventive Action(Filled by Maint)")
                    ),
                    sheets_after_sanding=_clean(
                        cell(row, "Problem observed after how many sheet  sanding")
                    ),
                    import_key=key,
                    reason="Already imported — write-up columns refreshed."
                    if action == "update"
                    else None,
                )
            )

        result.unknown_machines = sorted(unknown)
        return result
    finally:
        workbook.close()
        buffer.close()
